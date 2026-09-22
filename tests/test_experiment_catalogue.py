import csv
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
PRED = REPO / "outputs" / "predictions"
SUMMARY = REPO / "reports" / "fydp3_summary.csv"


@unittest.skipUnless(
    (PRED / "15_pipeline_qwen9b_BEST_RESULT.jsonl").exists() and SUMMARY.exists(),
    "result files not present",
)
class CatalogueAgreesWithTheThesisTests(unittest.TestCase):
    """The per-system benchmark must never contradict the numbers already reported.

    `fydp3_summary.csv` is what the dossier and the Pareto figure were built from.
    The catalogue recomputes every system independently, through different code, so
    agreement here is a real cross-check rather than a copy.
    """

    @classmethod
    def setUpClass(cls):
        from scripts.build_experiment_catalogue import load_data, measure, systems

        data = load_data()
        cls.by_slug = {}
        for s in systems():
            measure(s, data)
            cls.by_slug[s.slug] = s
        with SUMMARY.open(encoding="utf-8") as handle:
            cls.summary = {row["condition"]: row for row in csv.DictReader(handle)}

    def correct(self, prefix):
        (match,) = [s for slug, s in self.by_slug.items() if slug.startswith(prefix)]
        return match.metrics["correct"]

    def reported(self, condition):
        return int(self.summary[condition]["correct"])

    def test_free_systems_match(self):
        self.assertEqual(self.correct("02_"), self.reported("local only (1 sample)"))
        self.assertEqual(self.correct("03_"), self.reported("blind retry, take last"))
        self.assertEqual(self.correct("04_"), self.reported("self-consistency@3 (majority)"))

    def test_every_supervised_arm_matches(self):
        pairs = {
            "08_": "hint-none", "09_": "hint-short", "10_": "hint-full",
            "11_": "more-escalation", "12_": "smart-gate", "18_": "BEST-qwen9b",
            "13_": "hint-none + voting", "14_": "hint-short + voting",
            "15_": "hint-full + voting", "16_": "more-escalation + voting",
            "17_": "smart-gate + voting", "19_": "BEST-qwen9b + voting",
        }
        for prefix, condition in pairs.items():
            with self.subTest(prefix=prefix, condition=condition):
                self.assertEqual(self.correct(prefix), self.reported(condition))

    def test_checker_tokens_match(self):
        (best,) = [s for slug, s in self.by_slug.items() if slug.startswith("19_")]
        self.assertAlmostEqual(
            best.metrics["checker_tokens_per_q"],
            float(self.summary["BEST-qwen9b + voting"]["cloud_tokens_per_q"]),
            places=1,
        )

    def test_marking_only_runs_leave_the_score_unchanged(self):
        # They grade answers and retry nothing, so they must equal system 02 exactly.
        for prefix in ("05_", "06_", "07_"):
            with self.subTest(prefix=prefix):
                self.assertEqual(self.correct(prefix), self.correct("02_"))

    def test_folder_numbers_are_unique_and_sequential(self):
        numbers = sorted(int(slug.split("_", 1)[0]) for slug in self.by_slug)
        self.assertEqual(numbers, list(range(1, len(numbers) + 1)))




XLSX = REPO / "reports" / "benchmark_all_systems.xlsx"
XLSX_COPY = REPO / "docs" / "experiments" / "benchmark_all_systems.xlsx"


@unittest.skipUnless(XLSX.exists() and XLSX_COPY.exists(), "spreadsheet not present")
class SpreadsheetCopiesAgreeTests(unittest.TestCase):
    """The copy beside the experiment folders must say what the reports copy says.

    It went stale once: the folder copy still described experiment 07 as "a perfect
    marker" after that was corrected to "the answer key, no second model". Two
    spreadsheets with the same name and different contents is the worst outcome, so
    the generator now writes both and this pins them together.
    """

    def test_both_copies_hold_the_same_values(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl not installed")

        for sheet in ("Benchmark", "Column guide"):
            with self.subTest(sheet=sheet):
                first = [list(r) for r in load_workbook(XLSX)[sheet].iter_rows(values_only=True)]
                second = [list(r) for r in load_workbook(XLSX_COPY)[sheet].iter_rows(values_only=True)]
                self.assertEqual(first, second)


class ColumnNotesTests(unittest.TestCase):
    """Every column in the benchmark spreadsheet must explain itself.

    The spreadsheet is read by people who have not seen the code. A column added
    later without a note would be a bare abbreviation like `judge_seconds`, which is
    exactly what the notes exist to prevent -- so a missing note fails the build.
    """

    def test_every_column_has_a_note(self):
        from scripts.build_experiment_catalogue import COLUMN_NOTES, CSV_FIELDS

        missing = [field for field in CSV_FIELDS if field not in COLUMN_NOTES]
        self.assertEqual(missing, [])

    def test_no_note_describes_a_column_that_does_not_exist(self):
        from scripts.build_experiment_catalogue import COLUMN_NOTES, CSV_FIELDS

        stale = [field for field in COLUMN_NOTES if field not in CSV_FIELDS]
        self.assertEqual(stale, [])

    def test_each_note_has_a_readable_title_and_explanation(self):
        from scripts.build_experiment_catalogue import COLUMN_NOTES

        for field, (title, explanation) in COLUMN_NOTES.items():
            with self.subTest(field=field):
                self.assertNotIn("_", title, "the title is for people, not code")
                self.assertGreater(len(explanation), 20)


if __name__ == "__main__":
    unittest.main()
