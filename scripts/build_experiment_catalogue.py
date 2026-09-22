"""One folder per experiment: what the system is, two diagrams, and its numbers.

Writes, for every system this thesis measured end to end on the full test set:

    docs/experiments/<NN_system-type>/README.md
    docs/experiments/<NN_system-type>/architecture.svg / .png
    docs/experiments/<NN_system-type>/question_flow.svg / .png

plus the combined benchmark:

    docs/experiments/README.md              plain-language benchmark tables
    reports/benchmark_all_systems.csv       the same numbers, machine-readable

Every number is counted from the result files in outputs/predictions/ at build time,
including the counts printed inside the diagrams. Nothing is typed in by hand, so a
re-run of any experiment is picked up by re-running this script.

The one exception is time. No full run recorded its own duration, so time is composed
from component speeds that were measured separately, and is labelled as an estimate
wherever it appears. See `thesis_pipeline/benchmark.py`.

    python scripts/build_experiment_catalogue.py
    python scripts/build_experiment_catalogue.py --no-render     tables only, fast
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from dataclasses import dataclass, field
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

from analyze_supervision import stack_voting  # noqa: E402
from thesis_pipeline.benchmark import estimated_seconds, flips, judge_quality  # noqa: E402
from thesis_pipeline.gate import cluster_answers, majority_answer  # noqa: E402
from thesis_pipeline.gsm8k import answers_match, extract_final_answer  # noqa: E402

PRED = REPO / "outputs" / "predictions"
OUT = REPO / "docs" / "experiments"
CSV_OUT = REPO / "reports" / "benchmark_all_systems.csv"
LATENCY = REPO / "reports" / "latency_benchmark.json"
LOG = REPO / "reports" / "experiment_log.jsonl"

# GLM-4.7-Flash seconds per judgement: the mean over 729 requests in the llama-server
# log, recorded in THESIS_DOSSIER.md section 5.9. It is not in latency_benchmark.json,
# which timed the Qwen checker, so it is carried here with its source.
GLM_JUDGE_SECONDS = 2.07


# -- loading ------------------------------------------------------------------


def load(name: str) -> dict[int, dict]:
    path = PRED / f"{name}.jsonl"
    with path.open(encoding="utf-8") as handle:
        return {int(r["id"]): r for r in (json.loads(line) for line in handle if line.strip())}


def final_of(row: dict) -> str | None:
    return row.get("pred_final_answer") or extract_final_answer(row.get("prediction"))


@dataclass
class Data:
    base: dict[int, dict]
    try1: dict[int, dict]
    try2: dict[int, dict]
    try3: dict[int, dict]
    ids: list[int]
    solver_tps: float
    qwen_judge_seconds: float

    @property
    def bank(self) -> dict[int, list[str | None]]:
        return {i: [final_of(self.try2[i]), final_of(self.try3[i])] for i in self.ids}

    def gold(self, i: int) -> str:
        return self.try1[i]["gold_final_answer"]

    def reference_correct(self) -> list[bool]:
        """The fine-tuned model answering once: every flip is counted against it."""
        return [bool(self.try1[i]["correct"]) for i in self.ids]


def load_data() -> Data:
    latency = json.loads(LATENCY.read_text(encoding="utf-8"))
    checker = next(c for c in latency["components"] if c["what"].startswith("checker"))
    try1 = load("02_answers_finetuned_try1_main")
    return Data(
        base=load("01_answers_untrained_baseline"),
        try1=try1,
        try2=load("03_answers_finetuned_try2"),
        try3=load("04_answers_finetuned_try3"),
        ids=sorted(try1),
        solver_tps=float(latency["local_tokens_per_second"]),
        qwen_judge_seconds=float(checker["mean"]),
    )


# -- the systems --------------------------------------------------------------


@dataclass
class System:
    slug: str
    name: str                      # plain language, one line
    technical: str                 # the same system in standard terms
    what_it_does: str              # a short paragraph, plain language
    kind: str                      # single | blind | vote | marking | cascade
    source: str | None = None      # result file stem
    trained: bool = True
    stacked: bool = False
    checker: str | None = None     # glm | qwen | oracle
    hint: str | None = None        # none | short | full
    gate: str | None = None        # agreement | smart
    budget: str | None = None      # "30%" | "37.3%"
    note: str = ""
    metrics: dict = field(default_factory=dict)
    flow: dict = field(default_factory=dict)


CHECKER_NAME = {
    "glm": "a large AI marker (GLM-4.7-Flash, 30 billion)",
    "qwen": "a medium AI marker (Qwen3.5-9B, 9 billion)",
    "oracle": "the answer key itself, compared in code (no second model)",
}
HINT_TEXT = {
    "none": "only yes or no, with no explanation",
    "short": "yes or no, plus which step went wrong",
    "full": "yes or no, which step went wrong, and how the question should be read",
}


MARKER_REPLY = {
    "none": "it just says wrong",
    "short": "it says which step went wrong",
    "full": "it says which step went wrong<br/>and how to read the question",
}
RETRY_TEXT = {
    "none": "told only that it was wrong",
    "short": "told which step went wrong",
    "full": "told which step went wrong<br/>and how to read the question",
}


def systems() -> list[System]:
    glm_arms = [
        ("08", "no-hint", "none", "agreement", "30%", "10_pipeline_glm30b_hint_none"),
        ("09", "short-hint", "short", "agreement", "30%", "11_pipeline_glm30b_hint_short"),
        ("10", "full-hint", "full", "agreement", "30%", "12_pipeline_glm30b_hint_full"),
        ("11", "full-hint_37pct", "full", "agreement", "37.3%", "13_pipeline_glm30b_hint_full_more_escalation"),
        ("12", "full-hint", "full", "smart", "30%", "14_pipeline_glm30b_hint_full_smart_gate"),
    ]
    out = [
        System(
            "01_base-model-only",
            "The small AI before we trained it",
            "Untrained gemma-4-E2B-it, 8-shot chain-of-thought prompt",
            "The model exactly as Google released it, answering each question once. Because "
            "an untrained model does not know the answer format, it is shown 8 solved example "
            "problems before every question. This is the starting point everything else is "
            "compared against.",
            "single", "01_answers_untrained_baseline", trained=False,
        ),
        System(
            "02_finetuned-model-only",
            "The small AI after training, answering once",
            "QLoRA fine-tuned gemma-4-E2B-it, zero-shot greedy decoding",
            "The same small AI after we trained it on 7,500 school maths problems. It answers "
            "each question once, with no help and no second attempt.",
            "single", "02_answers_finetuned_try1_main",
        ),
        System(
            "03_finetuned-model+blind-retry",
            "Trained AI, asked again, keeping the second answer",
            "Fine-tuned model, resample and take the last sample (no verifier)",
            "The trained AI answers, then is simply asked again, and the second answer is kept. "
            "Nobody checks either answer. This is the control that shows whether a retry helps "
            "on its own, without any marker.",
            "blind",
        ),
        System(
            "04_finetuned-model+majority-vote",
            "Trained AI answers three times, most common answer kept",
            "Fine-tuned model with self-consistency@3 (majority vote)",
            "The trained AI answers each question three times and the most common answer is "
            "kept. No other AI is involved and nothing leaves the machine. This is the free "
            "alternative every system with a marker has to beat.",
            "vote",
        ),
        System(
            "05_finetuned-model+checker-marks-all-answers_glm-30b",
            "A large AI marks every answer (measuring the marker)",
            "Verdict-only pass: GLM-4.7-Flash judges every attempt-1 answer, no retry",
            "The large AI marker grades all 1,319 answers, but the small AI is never asked to "
            "try again, so no answer changes. This experiment does not produce a better score; "
            "it measures how good the marker is at spotting wrong answers.",
            "marking", "16_pipeline_glm30b_no_gate_every_question", checker="glm",
            note="The result file's name says 'pipeline ... every question', but every row has "
                 "exactly one attempt: it is a marking-only run, not a full system.",
        ),
        System(
            "06_finetuned-model+checker-marks-all-answers_qwen-9b",
            "A medium AI marks every answer (measuring the marker)",
            "Verdict-only pass: Qwen3.5-9B judges every attempt-1 answer, no retry",
            "The medium AI marker grades all 1,319 answers, with no retries, to measure how good "
            "it is at spotting wrong answers. Compare it with experiment 05.",
            "marking", "08_checker_qwen9b_marks_all", checker="qwen",
        ),
        System(
            "07_finetuned-model+perfect-checker-marks-all-answers",
            "The answer key used as the marker (the upper limit)",
            "Verdict-only pass with the exact-match oracle; no second model is loaded",
            "**Only one model runs in this experiment: the small AI.** There is no second model. The 'marker' is a few lines of code that compare each answer with the answer key, which is why it is never wrong. Nothing leaves the machine and no tokens are exchanged. It is here for two reasons: it shows the best score any marker could possibly reach, so the real markers in experiments 05 and 06 have something to be measured against, and it confirms our scoring code is correct — an answer key that scored anything but 100% would mean a bug.",
            "marking", "09_checker_perfect_oracle_marks_all", checker="oracle",
        ),
    ]
    for stacked in (False, True):
        for num, suffix, hint, gate, budget, source in glm_arms:
            n = int(num) + (5 if stacked else 0)
            gate_part = "smart-gate" if gate == "smart" else "gate"
            vote_part = "vote+" if stacked else ""
            out.append(System(
                f"{n:02d}_finetuned-model+{vote_part}{gate_part}+checker_glm-30b_{suffix}",
                _cascade_name(stacked, gate, budget, hint, "large"),
                _cascade_technical(stacked, gate, budget, hint, "GLM-4.7-Flash"),
                _cascade_blurb(stacked, gate, budget, hint, "large"),
                "cascade", source, stacked=stacked, checker="glm", hint=hint, gate=gate, budget=budget,
            ))
    for stacked, n in ((False, 18), (True, 19)):
        vote_part = "vote+" if stacked else ""
        out.append(System(
            f"{n}_finetuned-model+{vote_part}smart-gate+checker_qwen-9b_full-hint",
            _cascade_name(stacked, "smart", "30%", "full", "medium")
            + (" — BEST RESULT" if stacked else ""),
            _cascade_technical(stacked, "smart", "30%", "full", "Qwen3.5-9B"),
            _cascade_blurb(stacked, "smart", "30%", "full", "medium"),
            "cascade", "15_pipeline_qwen9b_BEST_RESULT", stacked=stacked, checker="qwen",
            hint="full", gate="smart", budget="30%",
        ))
    return out


def _cascade_name(stacked, gate, budget, hint, size):
    parts = ["Trained AI"]
    if stacked:
        parts.append("majority vote")
    parts.append(f"{'smart ' if gate == 'smart' else ''}selection of hard questions ({budget})")
    parts.append(f"{size} AI marker, {hint} hint")
    return " + ".join(parts)


def _cascade_technical(stacked, gate, budget, hint, checker):
    gate_text = "combined gate (disagreement, confidence tie-break)" if gate == "smart" else "disagreement gate"
    level = {"none": "L0", "short": "L1", "full": "L2"}[hint]
    return (f"{'Stacked on self-consistency@3; ' if stacked else ''}{gate_text} at {budget} "
            f"escalation; {checker} verifier; {level} feedback")


def _cascade_blurb(stacked, gate, budget, hint, size):
    choose = ("the ones where all three attempts disagreed, taking first the ones the AI "
              "sounded least sure about" if gate == "smart"
              else "the ones where the three attempts disagreed most")
    first = (
        "The trained AI answers every question three times. "
        + ("Where at least two attempts agree, that answer is kept straight away. " if stacked
           else "On questions that are not sent for help, its first answer is kept. ")
    )
    return (
        first
        + f"About {budget} of questions are sent to {size} AI marker — {choose}. "
        + f"If the marker says the working is wrong, it replies with {HINT_TEXT[hint]}, "
        + "and the small AI tries again. The marker is never shown the correct answer."
    )


# -- measuring each system ----------------------------------------------------


def measure(system: System, data: Data) -> None:
    ids, n = data.ids, len(data.ids)
    gen = prompt = gens = calls = cloud_in = cloud_out = sent = 0
    errors = 0
    correct: list[bool]
    flow: dict = {}

    if system.kind == "single":
        rows = data.base if not system.trained else data.try1
        correct = [bool(rows[i]["correct"]) for i in ids]
        gen = sum(rows[i]["tokens_generated"] for i in ids)
        prompt = sum(rows[i]["prompt_tokens"] for i in ids)
        gens = n

    elif system.kind == "blind":
        correct = [answers_match(final_of(data.try3[i]), data.gold(i)) for i in ids]
        gen = sum(data.try1[i]["tokens_generated"] + data.try3[i]["tokens_generated"] for i in ids)
        prompt = sum(data.try1[i]["prompt_tokens"] + data.try3[i]["prompt_tokens"] for i in ids)
        gens = 2 * n

    elif system.kind == "vote":
        correct = []
        groups = {1: 0, 2: 0, 3: 0}
        for i in ids:
            samples = [final_of(data.try1[i]), *data.bank[i]]
            groups[min(len(cluster_answers(samples)), 3)] += 1
            correct.append(answers_match(majority_answer(samples), data.gold(i)))
        for part in (data.try1, data.try2, data.try3):
            gen += sum(part[i]["tokens_generated"] for i in ids)
            prompt += sum(part[i]["prompt_tokens"] for i in ids)
        gens = 3 * n
        flow["groups"] = groups

    elif system.kind == "marking":
        rows = load(system.source)
        correct = [bool(rows[i]["correct"]) for i in ids]
        gen = sum(data.try1[i]["tokens_generated"] for i in ids)
        prompt = sum(data.try1[i]["prompt_tokens"] for i in ids)
        gens = n
        for i in ids:
            calls += int(rows[i].get("supervisor_calls") or 0)
            cloud_in += int(rows[i].get("supervisor_input_tokens") or 0)
            cloud_out += int(rows[i].get("supervisor_output_tokens") or 0)
            errors += bool(rows[i]["attempts"][0].get("supervisor_error"))
        sent = n
        verdicts = [rows[i]["attempts"][0].get("supervisor_accepted") for i in ids]
        truth = [bool(rows[i]["attempts"][0]["exact_correct"]) for i in ids]
        quality = judge_quality(verdicts, truth)
        flow.update(
            quality=quality,
            accepted=sum(1 for v in verdicts if v),
            rejected=sum(1 for v in verdicts if v is False),
            rejected_wrong=sum(1 for v, t in zip(verdicts, truth) if v is False and not t),
            rejected_right=sum(1 for v, t in zip(verdicts, truth) if v is False and t),
            wrong_total=sum(1 for t in truth if not t),
        )

    elif system.kind == "cascade":
        rows = load(system.source)
        ordered = [rows[i] for i in ids]
        scored = stack_voting(ordered, data.bank) if system.stacked else ordered
        correct = [bool(r["correct"]) for r in scored]
        first_accept = first_reject = retried_once = retried_twice = 0
        sent_correct = kept_correct = 0
        for original, final in zip(ordered, scored):
            i = int(original["id"])
            gen += int(original.get("total_model_tokens") or 0)
            prompt += int(original.get("total_prompt_tokens") or 0)
            # The gate cannot decide without the two extra attempts, so every
            # question pays for them whether or not it is escalated.
            gen += data.try2[i]["tokens_generated"] + data.try3[i]["tokens_generated"]
            prompt += data.try2[i]["prompt_tokens"] + data.try3[i]["prompt_tokens"]
            gens += int(original.get("attempt_count") or 1) + 2
            calls += int(original.get("supervisor_calls") or 0)
            cloud_in += int(original.get("supervisor_input_tokens") or 0)
            cloud_out += int(original.get("supervisor_output_tokens") or 0)
            errors += bool((original.get("attempts") or [{}])[0].get("supervisor_error"))
            if original.get("escalated"):
                sent += 1
                sent_correct += bool(final["correct"])
                if (original["attempts"][0].get("supervisor_accepted")):
                    first_accept += 1
                else:
                    first_reject += 1
                count = int(original.get("attempt_count") or 1)
                retried_once += count == 2
                retried_twice += count == 3
            else:
                kept_correct += bool(final["correct"])
        groups = {1: 0, 2: 0, 3: 0}
        for i in ids:
            groups[min(len(cluster_answers([final_of(data.try1[i]), *data.bank[i]])), 3)] += 1
        flow.update(
            groups=groups, sent=sent, kept=n - sent, first_accept=first_accept,
            first_reject=first_reject, retried_once=retried_once, retried_twice=retried_twice,
            sent_correct=sent_correct, kept_correct=kept_correct,
        )
    else:
        raise ValueError(f"unknown system kind {system.kind!r}")

    judge_seconds = {"glm": GLM_JUDGE_SECONDS, "qwen": data.qwen_judge_seconds,
                     "oracle": 0.0, None: 0.0}[system.checker]
    seconds = estimated_seconds(gen, calls, data.solver_tps, judge_seconds)
    fixed, broken = flips(data.reference_correct(), correct)
    right = sum(correct)
    system.metrics = {
        "correct": right,
        "total": n,
        "accuracy": right / n,
        "vs_finetuned_points": 100 * (right - sum(data.reference_correct())) / n,
        "fixed": fixed,
        "broken": broken,
        "attempts_per_q": gens / n,
        "local_generated_per_q": gen / n,
        "local_read_per_q": prompt / n,
        "sent_to_checker": sent,
        "sent_share": sent / n,
        "checker_calls": calls,
        "checker_tokens_per_q": (cloud_in + cloud_out) / n,
        "checker_in_per_q": cloud_in / n,
        "checker_out_per_q": cloud_out / n,
        "checker_errors": errors,
        "judge_seconds": judge_seconds,
        "solver_tps": data.solver_tps,
        "est_seconds_per_q": seconds / n,
        "est_seconds_total": seconds,
    }
    system.flow = {**flow, "correct": right, "total": n}


# -- diagrams -----------------------------------------------------------------


def architecture(s: System) -> str:
    solver = ("<b>The small AI, before training</b><br/>shown 8 solved example problems<br/>before every question"
              if not s.trained else
              "<b>The small AI, trained by us</b><br/>on 7,500 school maths problems")
    lines = ["flowchart TB",
             '    subgraph DEV["RUNS ON YOUR OWN MACHINE"]',
             "        direction TB",
             f'        SOLVER["{solver}"]']
    tail = []
    if s.kind == "single":
        lines.append('        ONE["Answers each question <b>once</b>"]')
        lines.append("        SOLVER --> ONE")
        tail.append('    ONE --> OUT["<b>The answer you see</b>"]')
    elif s.kind == "blind":
        lines.append('        TWO["Answers, then is asked <b>again</b><br/>nobody checks either answer"]')
        lines.append("        SOLVER --> TWO")
        tail.append('    TWO --> OUT["<b>The answer you see</b><br/>always the second attempt"]')
    elif s.kind == "vote":
        lines.append('        VOTE["<b>Compare three attempts</b><br/>keep the most common answer"]')
        lines.append("        SOLVER -->|three attempts| VOTE")
        tail.append('    VOTE --> OUT["<b>The answer you see</b>"]')
    elif s.kind == "marking":
        lines.append('        ONE["Answers each question <b>once</b>"]')
        lines.append("        SOLVER --> ONE")
        if s.checker == "oracle":
            # No second model exists in this experiment: the "marker" is a few
            # lines of code comparing with the answer key. Drawing it outside the
            # device, as the real markers are drawn, invents a model that was
            # never run.
            lines.append('        KEY["<b>The answer key</b><br/>a few lines of code compare each<br/>'
                         'answer with the correct one<br/><b>no second AI is involved</b>"]')
            lines.append("        ONE --> KEY")
    else:
        if s.stacked:
            lines.append('        VOTE["<b>Compare three attempts</b><br/>if two or three agree,<br/>keep that answer"]')
            lines.append("        SOLVER -->|three attempts| VOTE")
        gate_text = ("<b>Pick the hard questions</b><br/>all three attempts disagreed;<br/>least sure-sounding first"
                     if s.gate == "smart" else
                     "<b>Pick the hard questions</b><br/>the attempts disagreed most")
        lines.append(f'        GATE["{gate_text}<br/>(about {s.budget} of questions)"]')
        lines.append("        SOLVER -->|three attempts| GATE")
    lines.append("    end")

    if s.kind in ("marking", "cascade") and s.checker != "oracle":
        who = CHECKER_NAME[s.checker]
        lines += ['    subgraph OFF["NOT ON YOUR MACHINE"]', "        direction TB",
                  f'        CHECK["<b>Marker</b><br/>{who}<br/><b>never shown the correct answer</b>"]',
                  "    end"]
    if s.kind == "marking":
        origin = "KEY" if s.checker == "oracle" else "CHECK"
        if s.checker != "oracle":
            lines.append("    ONE -->|every answer| CHECK")
        lines.append(f'    {origin} --> SCORE["<b>A mark for every answer</b><br/>'
                     "the answer is never changed:<br/>this measures the marking, "
                     'not the system"]')
    elif s.kind == "cascade":
        lines.append("    GATE -->|hard questions only| CHECK")
        lines.append(f'    CHECK -->|"if wrong: {MARKER_REPLY[s.hint]}<br/>(never the answer)"| SOLVER')
        lines.append('    SOLVER --> OUT["<b>The answer you see</b><br/>always written by the small AI"]')
        if s.stacked:
            lines.append("    VOTE -->|agreed answers| OUT")
    lines += tail
    lines += ["    style DEV fill:#e8f4ea,stroke:#2d6a4f,stroke-width:2px"]
    if s.kind in ("marking", "cascade") and s.checker != "oracle":
        lines.append("    style OFF fill:#fdf0e6,stroke:#b5651d,stroke-width:2px")
    if s.checker == "oracle" and s.kind == "marking":
        lines.append("    style KEY fill:#eeeeee,stroke:#888,stroke-dasharray: 4 4")
    return "\n".join(lines)


def question_flow(s: System) -> str:
    f, m = s.flow, s.metrics
    total = f["total"]
    right, wrong = f["correct"], total - f["correct"]
    L = ["flowchart TD", f'    Q["<b>{total:,} maths questions</b>"]']
    end = f'    RES["<b>{right:,} right</b> ({100 * right / total:.1f}%)<br/>{wrong:,} wrong"]'

    if s.kind == "single":
        how = "shown 8 solved examples first,<br/>then answers once" if not s.trained else "answers once"
        L += [f'    Q --> A["The small AI {how}"]', "    A --> RES", end]
    elif s.kind == "blind":
        L += ['    Q --> A["The small AI answers"]',
              '    A --> B["It is asked the same question again<br/>nobody checks either answer"]',
              '    B --> C["The second answer is kept"]', "    C --> RES", end]
    elif s.kind == "vote":
        g = f["groups"]
        L += ['    Q --> A["The small AI answers each question <b>three times</b>"]',
              '    A --> D{"Did the attempts agree?"}',
              f'    D -->|"all three matched<br/><b>{g[1]}</b>"| K["Keep the answer they agreed on"]',
              f'    D -->|"two of three matched<br/><b>{g[2]}</b>"| K',
              f'    D -->|"all three differed<br/><b>{g[3]}</b>"| F["No most-common answer,<br/>so keep the first attempt"]',
              "    K --> RES", "    F --> RES", end,
              "    style K fill:#e8f4ea,stroke:#2d6a4f"]
    elif s.kind == "marking" and s.checker == "oracle":
        L += ['    Q --> A["The small AI answers each question once"]',
              f'    A --> M["<b>Each answer is compared with the answer key</b><br/>'
              'no second AI: this is a few lines of code"]',
              f'    M -->|"matches<br/><b>{f["accepted"]}</b>"| OK["Counted right"]',
              f'    M -->|"does not match<br/><b>{f["rejected"]}</b>"| NO["Counted wrong"]',
              '    OK --> SAME["<b>No answer is changed.</b><br/>This experiment shows the best score any<br/>'
              'marker could reach, and checks that our<br/>scoring code is correct."]',
              "    NO --> SAME", "    SAME --> RES", end,
              "    style M fill:#eeeeee,stroke:#888,stroke-dasharray: 4 4"]
    elif s.kind == "marking":
        L += ['    Q --> A["The small AI answers each question once"]',
              f'    A --> M["<b>The marker grades all {total:,} answers</b>"]',
              f'    M -->|"says right<br/><b>{f["accepted"]}</b>"| OK["Marked right"]',
              f'    M -->|"says wrong<br/><b>{f["rejected"]}</b>"| NO["Marked wrong"]',
              f'    NO --> NW["<b>{f["rejected_wrong"]}</b> really were wrong<br/>caught {f["rejected_wrong"]} of the {f["wrong_total"]} wrong answers"]',
              f'    NO --> NR["<b>{f["rejected_right"]}</b> were actually right<br/>a mistake by the marker"]',
              '    OK --> SAME["<b>No answer is changed</b><br/>the score stays the same"]',
              "    NW --> SAME", "    NR --> SAME", "    SAME --> RES", end,
              "    style M fill:#fdf0e6,stroke:#b5651d", "    style NR fill:#fdece9,stroke:#b03a2e"]
    else:
        g = f["groups"]
        if s.stacked:
            L += ['    Q --> A["The small AI answers each question <b>three times</b>"]',
                  '    A --> D{"Did the attempts agree?"}',
                  f'    D -->|"all or two matched<br/><b>{g[1] + g[2]}</b>"| K["<b>Keep the agreed answer</b><br/>nothing leaves the machine"]',
                  f'    D -->|"all three differed<br/><b>{g[3]}</b>"| H["<b>The hard ones</b><br/>order them, least sure first"]']
            leftover = f["kept"] - (g[1] + g[2])
            if leftover > 0:
                L.append(f'    H -->|"<b>{leftover}</b> — over the<br/>{s.budget} allowance"| KF["Keep the first attempt"]')
                L.append("    KF --> RES")
            src = "H"
        else:
            L += ['    Q --> A["The small AI answers each question <b>three times</b>"]',
                  ('    A --> H["<b>Pick the hard questions</b><br/>attempts disagreed; least sure first"]'
                   if s.gate == "smart" else
                   '    A --> H["<b>Pick the hard questions</b><br/>where the attempts disagreed most"]'),
                  f'    H -->|"not picked<br/><b>{f["kept"]}</b>"| K["<b>Keep the first attempt</b><br/>nothing leaves the machine"]']
            src = "H"
        L += [f'    {src} -->|"picked<br/><b>{f["sent"]}</b>"| M["<b>The marker checks the working</b><br/>never shown the correct answer"]',
              f'    M -->|"says right<br/><b>{f["first_accept"]}</b>"| KEEP2["Keep that answer"]',
              f'    M -->|"says wrong<br/><b>{f["first_reject"]}</b>"| R1["<b>The small AI tries again</b><br/>{RETRY_TEXT[s.hint]}"]',
              f'    R1 -->|"marker now says right<br/><b>{f["retried_once"]}</b>"| KEEP2',
              f'    R1 -->|"still wrong<br/><b>{f["retried_twice"]}</b>"| R2["One last try,<br/>kept without checking"]',
              "    K --> RES", "    KEEP2 --> RES", "    R2 --> RES", end,
              "    style K fill:#e8f4ea,stroke:#2d6a4f", "    style M fill:#fdf0e6,stroke:#b5651d"]
    L.append("    style RES fill:#e6eefc,stroke:#2b4c8c,stroke-width:2px")
    return "\n".join(L)


# -- writing ------------------------------------------------------------------


def hms(seconds: float) -> str:
    minutes = round(seconds / 60)
    return f"{minutes // 60} h {minutes % 60:02d} min"


def per_system_readme(s: System) -> str:
    m = s.metrics
    rows = [
        ("Right answers", f"**{m['correct']:,} of {m['total']:,} — {100 * m['accuracy']:.2f}%**"),
        ("Compared with the trained AI answering once", f"{m['vs_finetuned_points']:+.2f} percentage points"),
        ("Answers turned from wrong to right", f"{m['fixed']}"),
        ("Answers turned from right to wrong", f"{m['broken']}"),
        ("AI attempts per question", f"{m['attempts_per_q']:.2f}"),
        ("Words the small AI writes per question (tokens)", f"{m['local_generated_per_q']:.1f}"),
        ("Words the small AI reads per question (tokens)", f"{m['local_read_per_q']:.1f}"),
        ("Questions sent to the marker", f"{m['sent_to_checker']:,} ({100 * m['sent_share']:.1f}%)"),
        ("Times the marker was asked", f"{m['checker_calls']:,}"),
        ("Tokens exchanged with the marker per question", f"{m['checker_tokens_per_q']:.1f}  (in {m['checker_in_per_q']:.1f} / out {m['checker_out_per_q']:.1f})"),
        ("Marker replies that could not be read (counted as 'right')", f"{m['checker_errors']}"),
        ("Small AI writing speed (measured)", f"{m['solver_tps']:.1f} tokens/second"),
        ("Marker time per check (measured)", ("—" if s.checker is None else "instant (answer key)" if s.checker == "oracle" else f"{m['judge_seconds']:.2f} s")),
        ("**Time per question (estimated)**", f"{m['est_seconds_per_q']:.1f} s"),
        ("**Time for all 1,319 questions (estimated)**", hms(m["est_seconds_total"])),
    ]
    table = "\n".join(f"| {k} | {v} |" for k, v in rows)
    quality = ""
    if s.kind == "marking":
        q = s.flow["quality"]
        quality = (
            "\n## How good is this marker?\n\n"
            "| | |\n|---|---|\n"
            f"| Wrong answers it caught | **{100 * q['recall']:.1f}%** |\n"
            f"| Right answers it wrongly failed | **{100 * q['false_reject_rate']:.1f}%** |\n"
            f"| When it said 'wrong', how often it was correct | {100 * q['precision']:.1f}% |\n"
        )
    caveat = ""
    if not s.trained:
        caveat = ("\n> **The time estimate does not hold for this system.** The writing speed was "
                  "measured with the training add-on loaded, and this untrained model runs without it. "
                  "The experiment log shows the real run was much faster — see *How accurate are the "
                  "time estimates?* in [`../README.md`](../README.md).\n")
    if s.gate == "smart":
        caveat += ("\n> The smart selection also reads the first answer once more to measure how sure "
                   "the AI sounded. That extra read was not timed and is not in the estimate.\n")
    note = f"\n> **Note:** {s.note}\n" if s.note else ""
    source = f"`outputs/predictions/{s.source}.jsonl`" if s.source else (
        "computed from `outputs/predictions/02_`, `03_` and `04_`" )
    return f"""# {s.slug.split('_', 1)[0]} — {s.name}

**In standard terms:** {s.technical}

{s.what_it_does}
{note}
## What it is made of

![architecture](architecture.svg)

## What happens to the questions

![question flow](question_flow.svg)

## Results

| | |
|---|---|
{table}
{quality}{caveat}
**Where these numbers come from:** {source}. Every count, including the ones inside the
diagrams, is computed from that file by `scripts/build_experiment_catalogue.py`.
Time is estimated from measured speeds, because the runs did not record their own
duration — see the main table in [`../README.md`](../README.md).

<details><summary>Diagram source (edit the generator, not this)</summary>

```mermaid
{architecture(s)}
```

```mermaid
{question_flow(s)}
```

</details>
"""


CSV_FIELDS = [
    "folder", "system", "technical", "correct", "total", "accuracy", "vs_finetuned_points",
    "fixed", "broken", "attempts_per_q", "local_generated_per_q", "local_read_per_q",
    "sent_to_checker", "sent_share", "checker_calls", "checker_tokens_per_q",
    "checker_in_per_q", "checker_out_per_q", "checker_errors", "solver_tps", "judge_seconds",
    "est_seconds_per_q", "est_seconds_total",
]


# (readable title, what it means) for every column. The spreadsheet is read by people
# who have not seen the code, so each header carries its explanation as a cell note,
# and the same text is repeated on a "Column guide" sheet for printing.
COLUMN_NOTES = {
    "folder": ("Folder",
               "Which folder inside docs/experiments holds this system's diagrams and description."),
    "system": ("System",
               "What the system is, in plain words."),
    "technical": ("System (technical terms)",
                  "The same system described in standard terminology, for use in the paper."),
    "correct": ("Right answers",
                "How many of the test questions this system answered correctly. An answer counts as "
                "right only when its final number matches the answer key."),
    "total": ("Questions",
              "How many questions were tested. Always 1,319: the complete GSM8K test set, no sampling."),
    "accuracy": ("Score",
                 "Right answers divided by questions. Shown as a percentage."),
    "vs_finetuned_points": ("Gain over trained AI alone",
                            "Percentage points above (or, if negative, below) system 02 — the trained small AI "
                            "answering each question once, which scored 55.88%."),
    "fixed": ("Turned right",
              "Questions that system 02 got wrong and this system got right."),
    "broken": ("Turned wrong",
               "Questions that system 02 got right and this system got wrong. The score alone hides "
               "these: a system can gain answers and lose others at the same time."),
    "attempts_per_q": ("AI attempts per question",
                       "On average, how many times the small AI answered each question. Includes the extra "
                       "attempts used for voting or for spotting hard questions, and any retries."),
    "local_generated_per_q": ("Tokens written per question",
                              "On average, how many tokens the small AI wrote per question, across all its "
                              "attempts. A token is roughly three-quarters of a word. This is what mostly "
                              "decides how long a system takes."),
    "local_read_per_q": ("Tokens read per question",
                         "On average, how many tokens the small AI had to read before answering, across all "
                         "its attempts. Very high for system 01, which is shown 8 solved examples first."),
    "sent_to_checker": ("Questions sent to marker",
                        "How many questions the marker (the second, larger AI) looked at, at least once."),
    "sent_share": ("Share sent to marker",
                   "Questions sent to the marker as a share of all questions. The thesis aims to keep this "
                   "low, since each one would leave the device."),
    "checker_calls": ("Marker calls",
                      "Total number of times the marker was asked to check an answer. Can exceed the "
                      "questions sent, because a retried answer is checked again."),
    "checker_tokens_per_q": ("Marker tokens per question",
                             "Tokens sent to plus received from the marker, averaged over ALL 1,319 "
                             "questions, including ones never sent. This is the cost the thesis measures: "
                             "what would travel to an online service. (Here the marker ran on the same "
                             "machine, so nothing actually left it.)"),
    "checker_in_per_q": ("Tokens sent to marker per question",
                         "The part of 'Marker tokens per question' that goes to the marker: the question "
                         "plus the small AI's working."),
    "checker_out_per_q": ("Tokens back from marker per question",
                          "The part of 'Marker tokens per question' that comes back: the verdict and any hint."),
    "checker_errors": ("Unreadable marker replies",
                       "Marker replies to the first check that could not be read, for example because they "
                       "were cut off. By design these count as 'right', so a fault never invents a rejection."),
    "solver_tps": ("Small AI writing speed (tokens/second)",
                   "How fast the trained small AI writes, measured over 40 questions on an RTX 4080 SUPER. "
                   "The same for every system, because it is the same model. Used to estimate time."),
    "judge_seconds": ("Seconds per marker check",
                      "Measured time for the marker to check one answer: about 1.00 s for the medium marker "
                      "(Qwen 9B), 2.07 s for the large one (GLM 30B). Zero when there is no marker, or for "
                      "the perfect marker, which just compares with the answer key."),
    "est_seconds_per_q": ("Estimated seconds per question",
                          "ESTIMATED, not measured: tokens written divided by writing speed, plus marker calls "
                          "times seconds per check. Checked against real runs, it is within 2-3.5% for systems "
                          "built on the trained AI. It does NOT hold for system 01."),
    "est_seconds_total": ("Estimated seconds for all 1,319",
                          "The same estimate for the whole test set, answering from scratch. Divide by 3,600 "
                          "for hours. The experiments themselves ran faster, because they reused answers "
                          "already saved to disk."),
}

XLSX_OUT = REPO / "reports" / "benchmark_all_systems.xlsx"
# A second copy sits beside the experiment folders, because that is where anyone
# browsing the catalogue looks for it. It is written here rather than copied by hand:
# the hand-made copy silently went stale once already, and a spreadsheet that
# disagrees with the folder it sits in is worse than no spreadsheet.
XLSX_COPY = OUT / "benchmark_all_systems.xlsx"
PERCENT_COLUMNS = {"accuracy", "sent_share"}


def write_xlsx(all_systems: list[System]) -> None:
    """The benchmark as an Excel workbook, with every column explained.

    The CSV stays the machine-readable copy; CSV cannot carry notes, so this is the
    one to open in Excel. Row 1 uses readable titles, each with a hover note giving
    its meaning, and a second sheet repeats every explanation for printing.
    """
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = "Benchmark"
    header_fill = PatternFill("solid", fgColor="DDE7F3")

    for col, key in enumerate(CSV_FIELDS, start=1):
        title, meaning = COLUMN_NOTES[key]
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        note = Comment(f"{meaning}\n\n(column name in the CSV: {key})", "benchmark")
        note.width, note.height = 320, 150
        cell.comment = note
        sheet.column_dimensions[get_column_letter(col)].width = (
            46 if key in ("system", "technical", "folder") else 16
        )

    for row, s in enumerate(all_systems, start=2):
        values = {"folder": s.slug, "system": s.name, "technical": s.technical, **s.metrics}
        for col, key in enumerate(CSV_FIELDS, start=1):
            cell = sheet.cell(row=row, column=col, value=values[key])
            if key in PERCENT_COLUMNS:
                cell.number_format = "0.00%"
            elif isinstance(values[key], float):
                cell.number_format = "0.00"
            if "BEST" in s.name:
                cell.font = Font(bold=True)

    sheet.row_dimensions[1].height = 45
    sheet.freeze_panes = "D2"  # keep the header and the system names in view

    guide = book.create_sheet("Column guide")
    guide.append(["Read me first"])
    guide.append(["Hover over any header on the Benchmark sheet to see its note. The same notes are listed "
                  "below. Every number is counted from the result files except the two time columns, which "
                  "are estimates."])
    guide.append([])
    guide.append(["Column", "What it means", "Column name in the CSV"])
    for key in CSV_FIELDS:
        title, meaning = COLUMN_NOTES[key]
        guide.append([title, meaning, key])
    guide["A1"].font = Font(bold=True, size=13)
    for cell in guide[4]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    guide.column_dimensions["A"].width = 36
    guide.column_dimensions["B"].width = 100
    guide.column_dimensions["C"].width = 26
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for destination in (XLSX_OUT, XLSX_COPY):
        destination.parent.mkdir(parents=True, exist_ok=True)
        book.save(destination)


def write_csv(all_systems: list[System]) -> None:
    CSV_OUT.parent.mkdir(parents=True, exist_ok=True)
    with CSV_OUT.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_FIELDS)
        writer.writeheader()
        for s in all_systems:
            m = s.metrics
            writer.writerow({
                "folder": s.slug, "system": s.name, "technical": s.technical,
                **{k: (f"{m[k]:.4f}" if isinstance(m[k], float) else m[k])
                   for k in CSV_FIELDS if k in m},
            })


def index_readme(all_systems: list[System], data: Data, ceiling: tuple[int, int]) -> str:
    def link(s):
        return f"[{s.slug.split('_', 1)[0]}]({s.slug}/README.md)"

    acc = ["| # | System | Right answers | Score | vs. trained AI alone | Turned right | Turned wrong |",
           "|---|---|---|---|---|---|---|"]
    for s in all_systems:
        m = s.metrics
        bold = "**" if "BEST" in s.name else ""
        acc.append(f"| {link(s)} | {bold}{s.name}{bold} | {m['correct']:,} | {bold}{100 * m['accuracy']:.2f}%{bold} | "
                   f"{m['vs_finetuned_points']:+.2f} | {m['fixed']} | {m['broken']} |")
    c, n = ceiling
    acc.append(f"| — | *Upper limit: counted right if any of the three attempts was right* | {c:,} | *{100 * c / n:.2f}%* | "
               f"*{100 * (c - sum(data.reference_correct())) / n:+.2f}* | — | — |")

    cost = ["| # | Attempts per question | Tokens written per question | Tokens read per question | "
            "Sent to marker | Marker calls | Marker tokens per question | Time per question* | Time for all 1,319* |",
            "|---|---|---|---|---|---|---|---|---|"]
    for s in all_systems:
        m = s.metrics
        dagger = " †" if not s.trained else ""
        cost.append(f"| {link(s)} | {m['attempts_per_q']:.2f} | {m['local_generated_per_q']:.1f} | "
                    f"{m['local_read_per_q']:.1f} | {m['sent_to_checker']:,} ({100 * m['sent_share']:.0f}%) | "
                    f"{m['checker_calls']:,} | {m['checker_tokens_per_q']:.1f} | "
                    f"{m['est_seconds_per_q']:.1f} s{dagger} | {hms(m['est_seconds_total'])}{dagger} |")

    markers = ["| # | Marker | Wrong answers caught | Right answers wrongly failed | Correct when it said 'wrong' | Time per check |",
               "|---|---|---|---|---|---|"]
    for s in all_systems:
        if s.kind != "marking":
            continue
        q = s.flow["quality"]
        t = "instant" if s.checker == "oracle" else f"{s.metrics['judge_seconds']:.2f} s"
        markers.append(f"| {link(s)} | {CHECKER_NAME[s.checker]} | {100 * q['recall']:.1f}% | "
                       f"{100 * q['false_reject_rate']:.1f}% | {100 * q['precision']:.1f}% | {t} |")

    folders = "\n".join(f"- [`{s.slug}/`]({s.slug}/README.md) — {s.name}" for s in all_systems)
    checks = timing_checks(data)
    timing_table = "\n".join(
        ["| Job | Estimated | Real (from the log) | Difference |", "|---|---|---|---|"]
        + [f"| {label} | {hms(est)} | {hms(real)} | "
           + (f"**{100 * (est - real) / real:+.0f}% — does not hold**" if abs(est - real) / real > 0.15
              else f"{100 * (est - real) / real:+.1f}%") + " |"
           for label, est, real in checks]
    )
    return f"""# Every experiment, one folder each

**Who this is for:** anyone writing the paper. Each folder below holds one system we
tested on all 1,319 test questions: a plain description, a diagram of what it is made
of, a diagram of what happens to the questions (with the real counts), and its numbers.

> **Generated file — do not edit by hand.** Every result here is counted from the
> result files by `scripts/build_experiment_catalogue.py`; the few measurements that
> were never saved to a file (the large marker's speed, the 30-question self-check)
> are quoted with their source. To update, re-run the script.

## The folders

Folder names describe the system: `+` joins the parts it is built from. For example
`finetuned-model+vote+smart-gate+checker_qwen-9b_full-hint` means *the trained AI, with
a majority vote, a smart way of picking hard questions, and a 9-billion-parameter
marker that gives a full hint*.

{folders}

---

## Table 1 — How many questions each system got right

"Turned right" and "turned wrong" are counted against the trained AI answering once
(system 02). A system can gain answers and lose others at the same time; the score
alone hides that.

{chr(10).join(acc)}

---

## Table 2 — What each system costs

{chr(10).join(cost)}

\\* **Time is estimated, not measured.** None of the full runs recorded how long it took.
Time is built from speeds that *were* measured: the small AI writes
**{data.solver_tps:.1f} tokens per second**, the medium marker takes
**{data.qwen_judge_seconds:.2f} s** per check, and the large marker **{GLM_JUDGE_SECONDS:.2f} s**
(sources below). Estimated time = tokens written ÷ writing speed + number of checks ×
time per check. Measured on an RTX 4080 SUPER; ordinary hardware would be slower.

† This estimate does not hold. The writing speed was measured with the training add-on
loaded; the untrained model runs without it, and its real run took about
{hms(checks[-1][2])} (table below).

### How accurate are the time estimates?

Three jobs ran back to back from a script, so the gap between their log entries is
close to how long they really took. Each is compared with the same estimation method,
applied to what that job actually had to do (earlier answers were reused from disk):

{timing_table}

So the estimates are reliable to within a few percent for every system built on the
trained AI, and should not be used for system 01.

### Two different "times"

The column above is how long each **system** would take to answer all 1,319 questions
from scratch. That is the number to compare systems by. The **experiments** themselves
ran faster, because they reused answers already saved to disk: systems 13–17 and 19 are
not separate runs at all, but the runs of 08–12 and 18 scored a second way.

**About "marker tokens".** In this project the marker ran on the same machine, so
nothing actually left it. The column counts the tokens that *would* travel to an
outside service if the marker were one — which is the cost the thesis argues about.

---

## Table 3 — How good each marker is

These three experiments only grade answers; nothing is retried, so they do not change
the score. They measure the marking itself.

**Experiment 07 is not a marker you could use.** Only one model runs in it — the small
AI. There is no second model: the "marker" is a few lines of code comparing each answer
with the answer key, so it is never wrong and nothing leaves the machine. Its "marker
calls" in Table 2 are local comparisons and cost no tokens. It is in this table to show
the best score any marker could reach, and to confirm our scoring code is correct.

{chr(10).join(markers)}

The 2-billion-parameter AI checking its own work was only tested on 30 questions, so it
has no folder here: it caught 22% of wrong answers and wrongly failed 42% of right ones
(`docs/THESIS_DOSSIER.md` §5.5.1).

---

## Measured speeds these estimates use

| What | Speed | Measured on | Source |
|---|---|---|---|
| Trained small AI writing an answer | {data.solver_tps:.2f} tokens/s | 40 questions | `reports/latency_benchmark.json` |
| Medium marker (Qwen3.5-9B) checking one answer | {data.qwen_judge_seconds:.2f} s | 40 checks | `reports/latency_benchmark.json` |
| Large marker (GLM-4.7-Flash) checking one answer | {GLM_JUDGE_SECONDS:.2f} s (44.6 tokens/s) | 729 checks | llama-server log, `docs/THESIS_DOSSIER.md` §5.9 |
| Training the small AI | 4,080 s (68 min) | the full training run | `reports/experiment_log.jsonl` |

## What is deliberately not in these tables

- **The 30-question self-check** (above) — too small to stand beside full runs.
- **The Qwen2.5-1.5B solver experiment** — stopped after 51 of 1,319 questions and not
  resumed, so it has no result (`docs/THESIS_DOSSIER.md` §5.10).
- **Smoke tests** of 3–20 questions, run only to check the code worked.
"""


def timing_checks(data: Data) -> list[tuple[str, float, float]]:
    """(label, estimated seconds, real seconds) for jobs whose duration the log reveals.

    The log records when a job finished, not when it started. The gap to the entry
    before is a real duration only when jobs ran back to back, so only such jobs are
    used. Each is estimated for what it actually did: the cascade arms reused saved
    answers and earlier checker verdicts, so they paid only for retries and new checks.
    """
    from datetime import datetime

    events = [json.loads(line) for line in LOG.read_text(encoding="utf-8").splitlines() if line.strip()]

    def gap_before(match) -> float:
        index = next(i for i, e in enumerate(events) if match(e))
        stamp = lambda e: datetime.fromisoformat(e["timestamp_utc"].replace("Z", "+00:00"))
        return (stamp(events[index]) - stamp(events[index - 1])).total_seconds()

    def command(e) -> str:
        return str(e.get("details", {}).get("command", ""))

    def arm_estimate(stem: str) -> float:
        rows = load(stem).values()
        retry = sum(int(r["total_model_tokens"]) - data.try1[int(r["id"])]["tokens_generated"] for r in rows)
        new = sum(int(r.get("supervisor_new_calls") or 0) for r in rows)
        return estimated_seconds(retry, new, data.solver_tps, GLM_JUDGE_SECONDS)

    tps = data.solver_tps
    return [
        ("One full set of 1,319 answers (second sample run)",
         sum(data.try3[i]["tokens_generated"] for i in data.ids) / tps,
         gap_before(lambda e: "samplebank_s2" in command(e))),
        ("System 09 as run: retries and new checks only",
         arm_estimate("11_pipeline_glm30b_hint_short"),
         gap_before(lambda e: "disagreement @ 30% / L1" in e["title"] and e["details"].get("examples") == 1319)),
        ("System 10 as run: retries and new checks only",
         arm_estimate("12_pipeline_glm30b_hint_full"),
         gap_before(lambda e: "disagreement @ 30% / L2" in e["title"])),
        ("System 01, the untrained model",
         sum(data.base[i]["tokens_generated"] for i in data.ids) / tps,
         gap_before(lambda e: "--few-shot 8" in command(e) and "--limit" not in command(e))),
    ]


def render_diagrams(folder: Path, sources: dict[str, str]) -> None:
    from scripts.render_diagrams import mermaid_command, render

    base = mermaid_command()
    scratch = folder / "_mermaid_src"
    scratch.mkdir(exist_ok=True)
    for stem, text in sources.items():
        src = scratch / f"{stem}.mmd"
        src.write_text(text + "\n", encoding="utf-8")
        render(base, src, folder / f"{stem}.svg", scale=None)
        render(base, src, folder / f"{stem}.png", scale=3)


def pass_at_3(data: Data) -> tuple[int, int]:
    hits = sum(
        any(answers_match(a, data.gold(i)) for a in [final_of(data.try1[i]), *data.bank[i]])
        for i in data.ids
    )
    return hits, len(data.ids)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--no-render", action="store_true", help="Write tables and READMEs; skip images.")
    parser.add_argument("--only", default=None, help="Render only folders whose name contains this.")
    args = parser.parse_args()

    data = load_data()
    all_systems = systems()
    for s in all_systems:
        measure(s, data)

    OUT.mkdir(parents=True, exist_ok=True)
    for s in all_systems:
        folder = OUT / s.slug
        folder.mkdir(parents=True, exist_ok=True)
        (folder / "README.md").write_text(per_system_readme(s), encoding="utf-8")
        rendered = ""
        if not args.no_render and (args.only is None or args.only in s.slug):
            render_diagrams(folder, {"architecture": architecture(s), "question_flow": question_flow(s)})
            rendered = "  + diagrams"
        m = s.metrics
        print(f"{s.slug:<66} {m['correct']:>4}  {100 * m['accuracy']:6.2f}%  "
              f"{m['checker_tokens_per_q']:6.1f} tok/q  ~{hms(m['est_seconds_total'])}{rendered}")

    (OUT / "README.md").write_text(index_readme(all_systems, data, pass_at_3(data)), encoding="utf-8")
    write_csv(all_systems)
    write_xlsx(all_systems)
    print(f"\nWrote {OUT / 'README.md'}\nWrote {CSV_OUT}\nWrote {XLSX_OUT}\nWrote {XLSX_COPY}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
