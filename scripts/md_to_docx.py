"""Convert the report Markdown into a .docx that opens cleanly in Google Docs.

Pandoc is not available on this machine, and the report is the kind of document
people will want to read in Word or Google Docs rather than on GitHub. This is a
small converter for exactly the Markdown the report uses -- ATX headings, pipe
tables, images, fenced code, blockquotes, bullet lists, horizontal rules and
inline bold/italic/code. It is not a general Markdown implementation and does not
try to be.

Two things are done deliberately for Google Docs rather than for Word:

* Table and cell widths are set in absolute units. Percentage widths survive Word
  but collapse in Google Docs.
* Every image is scaled to fit inside the text frame before it is inserted. The
  diagrams range from 2352x486 to 2352x4269 pixels, so inserting them at native
  size would push most of them off the page.

Usage:
    python scripts/md_to_docx.py docs/PAPER.md docs/PAPER.docx
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

# US Letter with one-inch margins.
PAGE_WIDTH = Inches(8.5)
MARGIN = Inches(1.0)
TEXT_WIDTH = Inches(6.5)
# Leaves room for a caption under a full-height figure rather than orphaning it.
MAX_IMAGE_HEIGHT = Inches(7.5)

# level -> (point size, colour, space before, space after).
HEADING_STYLE = {
    1: (20, RGBColor(0x11, 0x2B, 0x4F), 20, 10),
    2: (15, RGBColor(0x1A, 0x1A, 0x1A), 16, 6),
    3: (12.5, RGBColor(0x2B, 0x4C, 0x8C), 13, 4),
    4: (11.5, RGBColor(0x33, 0x33, 0x33), 10, 3),
}

INLINE = re.compile(
    r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*)|`[^`]+`)",
    re.S,
)


def add_inline(paragraph, text: str) -> None:
    """Write `text` into `paragraph`, honouring bold, italic and inline code.

    Splitting on a single alternation keeps the markers in document order, which
    a sequence of str.replace calls would not.
    """
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("***") and piece.endswith("***"):
            run = paragraph.add_run(piece[3:-3])
            run.bold = run.italic = True
        elif piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*"):
            paragraph.add_run(piece[1:-1]).italic = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x2E)
        else:
            paragraph.add_run(piece)


def fitted_size(path: Path) -> dict:
    """Width/height keywords that fit the image inside the text frame.

    Scaled on whichever dimension binds first, so aspect ratio is preserved for
    both the very wide and the very tall diagrams.
    """
    width_px, height_px = Image.open(path).size
    scale = min(TEXT_WIDTH / width_px, MAX_IMAGE_HEIGHT / height_px, 1.0 * 914400 / 96 / 1)
    # Work in EMU: start from 96 dpi, then shrink to fit.
    width_emu = width_px * 914400 // 96
    height_emu = height_px * 914400 // 96
    shrink = min(TEXT_WIDTH / width_emu, MAX_IMAGE_HEIGHT / height_emu, 1.0)
    return {"width": int(width_emu * shrink), "height": int(height_emu * shrink)}


def shade(cell, colour: str) -> None:
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")  # "solid" renders black
    element.set(qn("w:fill"), colour)
    cell._tc.get_or_add_tcPr().append(element)


def add_rule(document) -> None:
    """A horizontal rule as a paragraph bottom border, not an empty table."""
    paragraph = document.add_paragraph()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "BBBBBB")
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(document, rows: list[list[str]]) -> None:
    # Many result blocks in the report are written as headerless two-column
    # tables (`| | |`). Keeping that empty row renders as a blank shaded bar.
    has_header = any(cell.strip() for cell in rows[0])
    if not has_header:
        rows = rows[1:]
    columns = max(len(r) for r in rows)
    table = document.add_table(rows=0, cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Google Docs needs an explicit width on the table and on every cell.
    column_width = int(TEXT_WIDTH / columns)
    table.autofit = False
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for position in range(columns):
            cell = cells[position]
            cell.width = column_width
            cell.paragraphs[0].text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            text = row[position] if position < len(row) else ""
            add_inline(paragraph, text)
            if index == 0 and has_header:
                shade(cell, "DDE7F3")
                for run in paragraph.runs:
                    run.bold = True


def convert(source: Path, destination: Path) -> int:
    document = Document()

    section = document.sections[0]
    section.page_width = PAGE_WIDTH
    section.page_height = Inches(11)
    section.left_margin = section.right_margin = MARGIN
    section.top_margin = section.bottom_margin = Inches(0.9)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    lines = source.read_text(encoding="utf-8").splitlines()
    base = source.parent
    index = 0
    images = 0
    tables = 0
    seen_first_heading = False

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        # Fenced code
        if stripped.startswith("```"):
            index += 1
            body = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                body.append(lines[index])
                index += 1
            index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.space_before = Pt(6)
            run = paragraph.add_run("\n".join(body))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue

        # Horizontal rule
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            add_rule(document)
            index += 1
            continue

        # Image on its own line
        image = re.fullmatch(r"!\[([^\]]*)\]\(([^\)]+)\)", stripped)
        if image:
            path = (base / image.group(2)).resolve()
            if path.exists():
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(path), **fitted_size(path))
                images += 1
            else:
                print(f"  missing image: {image.group(2)}", file=sys.stderr)
            index += 1
            continue

        # Heading
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 4)
            text = heading.group(2).strip()
            paragraph = document.add_heading("", level=level)
            add_inline(paragraph, text)

            # Word's built-in heading sizes leave level 3 barely distinguishable
            # from bold body text, and every system in Chapter 8 is a level 3.
            size, colour, before, after = HEADING_STYLE[level]
            for run in paragraph.runs:
                run.font.size = Pt(size)
                run.font.color.rgb = colour
                run.bold = True
            paragraph.paragraph_format.space_before = Pt(before)
            paragraph.paragraph_format.space_after = Pt(after)

            if level == 1:
                if seen_first_heading:
                    # Each chapter and appendix opens a page, as a report should.
                    paragraph.insert_paragraph_before().add_run().add_break(
                        WD_BREAK.PAGE
                    )
                else:
                    # The title page block.
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size = Pt(26)
                seen_first_heading = True
            index += 1
            continue

        # Pipe table
        if stripped.startswith("|") and index + 1 < len(lines) and re.fullmatch(
            r"\|[\s:|-]+\|", lines[index + 1].strip()
        ):
            rows = [split_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_row(lines[index]))
                index += 1
            add_table(document, rows)
            document.add_paragraph()
            tables += 1
            continue

        # Blockquote
        if stripped.startswith(">"):
            body = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                body.append(lines[index].strip().lstrip(">").strip())
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.35)
            add_inline(paragraph, " ".join(body))
            for run in paragraph.runs:
                run.italic = True
            continue

        # Bullet or numbered list item, with its continuation lines
        bullet = re.match(r"^\s*([-*+]|\d+\.)\s+(.*)$", line)
        if bullet:
            ordered = bullet.group(1)[0].isdigit()
            body = [bullet.group(2)]
            index += 1
            while index < len(lines) and lines[index].startswith(("   ", "\t")) \
                    and lines[index].strip() and not re.match(r"^\s*([-*+]|\d+\.)\s", lines[index]):
                body.append(lines[index].strip())
                index += 1
            style = "List Number" if ordered else "List Bullet"
            paragraph = document.add_paragraph(style=style)
            add_inline(paragraph, " ".join(body))
            continue

        # Plain paragraph: join wrapped lines until a blank or a block starter
        body = [stripped]
        index += 1
        while index < len(lines):
            nxt = lines[index]
            if not nxt.strip() or nxt.strip().startswith(("#", "|", ">", "```", "![")) \
                    or re.fullmatch(r"-{3,}", nxt.strip()) \
                    or re.match(r"^\s*([-*+]|\d+\.)\s", nxt):
                break
            body.append(nxt.strip())
            index += 1
        add_inline(document.add_paragraph(), " ".join(body))

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    print(f"wrote {destination}  ({images} images, {tables} tables)")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    return convert(args.source, args.destination)


if __name__ == "__main__":
    raise SystemExit(main())
